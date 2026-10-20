# -*- coding: utf-8 -*-
"""Pruebas del traductor de errores y reintentos."""

import openpyxl
import pytest

import errores as err


def test_traduce_permission_denied():
    exc = PermissionError(13, "Permission denied")
    mensaje = err.traducir_error_excel(exc)
    assert "abierto en Excel o bloqueado por OneDrive" in mensaje
    assert "SIMULAR" in mensaje


def test_traduce_archivo_inexistente(tmp_path):
    mapeo, mensaje = err.leer_mapeo_con_reintentos(
        tmp_path / "no_existe.xlsx", intentos=1
    )
    assert mapeo is None
    assert "ya no existe" in mensaje


def test_traduce_no_excel(tmp_path):
    ruta = tmp_path / "archivo.txt"
    ruta.write_text("x", encoding="utf-8")
    mapeo, mensaje = err.leer_mapeo_con_reintentos(ruta, intentos=1)
    assert mapeo is None
    assert "no es un Excel valido" in mensaje


def test_excel_vacio_devuelve_dict_vacio(tmp_path):
    ruta = tmp_path / "plantilla.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "ID Origen"
    wb.save(ruta)

    mapeo, mensaje = err.leer_mapeo_con_reintentos(ruta, intentos=1)
    assert mapeo == {}
    assert mensaje == ""


def test_excel_columnas_invertidas(tmp_path):
    ruta = tmp_path / "m.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["9900000001", "2026XX840820X1"])
    wb.save(ruta)

    mapeo, mensaje = err.leer_mapeo_con_reintentos(ruta, intentos=1)
    assert mensaje == ""
    assert mapeo == {"2026XX840820X1": "9900000001"}


def test_reintentos_ante_bloqueo(monkeypatch, tmp_path):
    ruta = tmp_path / "m.xlsx"
    ruta.write_bytes(b"x")
    intentos = {"n": 0}
    original = err._cargar_mapeo_interno

    def fallar_dos_veces(*args, **kwargs):
        intentos["n"] += 1
        if intentos["n"] < 3:
            raise PermissionError(13, "Permission denied")
        return {"2026XX840820X1": "9900000001"}

    monkeypatch.setattr(err, "_cargar_mapeo_interno", fallar_dos_veces)
    mapeo, mensaje = err.leer_mapeo_con_reintentos(ruta, intentos=3, espera=0)

    assert intentos["n"] == 3
    assert mensaje == ""
    assert mapeo == {"2026XX840820X1": "9900000001"}
    monkeypatch.setattr(err, "_cargar_mapeo_interno", original)