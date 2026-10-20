# -*- coding: utf-8 -*-
"""
Pruebas de deteccion universal: nombres aleatorios, estructuras variadas,
destino alfanumerico, modo personalizado y fuzz.
"""

import random
import string

import openpyxl
import pytest

import renamer_core as core

PDF_OK = b"%PDF-1.4 contenido de prueba"

NOMBRES_ALEATORIOS = [
    "INV-2026-0001 factura.pdf",
    "documento_contrato_final.pdf",
    "expediente_2026_001.pdf",
    "RECIBO-4567-acuse.pdf",
    "2026EE963171O1 radicado tipico.pdf",
    "acta de reunion sala 3.pdf",
    "12345678.pdf",
    "OF-2026-777 anexo 2.pdf",
    "certificado_medico_02.pdf",
    "2026XX840820X1 ACME CORPORACION.pdf",
]


@pytest.fixture(autouse=True)
def codigo_limpio():
    core.CODIGO_FIJO = ""
    yield
    core.CODIGO_FIJO = ""


def crear_pdf(carpeta, nombre, contenido=PDF_OK):
    (carpeta / nombre).write_bytes(contenido)


def crear_excel(ruta, filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Titulos", "Libres"])
    for fila in filas:
        ws.append(fila)
    wb.save(ruta)


# --- Origen universal ----------------------------------------------------------


def test_origenes_aleatorios_todos_reconocidos(tmp_path):
    origen = tmp_path / "origen"
    origen.mkdir()
    for nombre in NOMBRES_ALEATORIOS:
        crear_pdf(origen, nombre)

    regs = core.inventariar(origen, modo=core.MODO_AUTO)

    assert core.listar_anomalias(regs) == []
    ids = {r["archivo"]: r["id_origen"] for r in regs}
    assert ids["2026EE963171O1 radicado tipico.pdf"] == "2026EE963171O1"
    assert ids["2026XX840820X1 ACME CORPORACION.pdf"] == "2026XX840820X1"
    assert ids["INV-2026-0001 factura.pdf"] == "INV-2026-0001"
    assert ids["documento_contrato_final.pdf"] == "documento_contrato_final"
    assert ids["12345678.pdf"] == "12345678"
    assert ids["certificado_medico_02.pdf"] == "certificado_medico_02"


def test_fuzz_cincuenta_nombres(tmp_path):
    origen = tmp_path / "origen"
    origen.mkdir()
    azar = random.Random(42)
    nombres = set()
    while len(nombres) < 50:
        piezas = [
            "".join(azar.choices(string.ascii_uppercase + string.digits,
                                 k=azar.randint(3, 8)))
            for _ in range(azar.randint(1, 3))
        ]
        separador = azar.choice([" ", "_", "-", ""])
        nombre = separador.join(piezas).strip("._- ") or "doc"
        if " " in nombre:
            nombre = nombre + " desc"
        nombres.add(nombre + ".pdf")

    for nombre in nombres:
        crear_pdf(origen, nombre)

    regs = core.inventariar(origen, modo=core.MODO_AUTO)

    assert core.listar_anomalias(regs) == []
    detectados = {r["id_origen"] for r in regs}
    assert len(detectados) >= 25


def test_modo_prefijo_con_espacios_internos(tmp_path):
    origen = tmp_path / "origen"
    origen.mkdir()
    crear_pdf(origen, "OF-2026-777 anexo adicional.pdf")
    regs = core.inventariar(origen, modo=core.MODO_AUTO)
    assert regs[0]["id_origen"] == "OF-2026-777"


# --- Excel universal -----------------------------------------------------------


def test_excel_destino_alfanumerico(tmp_path):
    ruta = tmp_path / "m.xlsx"
    crear_excel(ruta, [("2026EE963171O1 radicado 1", "EXP-2026-001")])
    mapeo, conflictos, informe = core.cargar_mapeo_detallado(
        ruta, modo=core.MODO_AUTO
    )

    assert mapeo == {"2026EE963171O1": "EXP-2026-001"}
    assert informe["pares"] == 1
    assert conflictos == {}


def test_excel_prefijo_con_descripcion(tmp_path):
    ruta = tmp_path / "m.xlsx"
    crear_excel(ruta, [("INV-2026-0001 factura", "9900000001")])
    mapeo, _conflictos, _informe = core.cargar_mapeo_detallado(
        ruta, modo=core.MODO_AUTO
    )
    assert mapeo == {"INV-2026-0001": "9900000001"}


def test_excel_ignora_encabezados(tmp_path):
    ruta = tmp_path / "m.xlsx"
    crear_excel(
        ruta,
        [("Radicado", "Expediente"), ("2026EE963171O1", "9900000001")],
    )
    mapeo, _conflictos, _informe = core.cargar_mapeo_detallado(
        ruta, modo=core.MODO_AUTO
    )
    assert "Radicado" not in mapeo
    assert mapeo == {"2026EE963171O1": "9900000001"}


def test_excel_fila_incompleta_se_reporta(tmp_path):
    ruta = tmp_path / "m.xlsx"
    crear_excel(
        ruta,
        [
            ("INV-2026-0001 factura", "9900000001"),
            ("INV-2026-0002 factura", None),
        ],
    )
    mapeo, _conflictos, informe = core.cargar_mapeo_detallado(
        ruta, modo=core.MODO_AUTO
    )
    assert len(mapeo) == 1
    saltadas = [motivo for _fila, motivo in informe["saltadas"]]
    assert any("destino" in motivo for motivo in saltadas)


# --- Flujo completo universal -----------------------------------------------------


def test_flujo_universal_completo(tmp_path):
    origen = tmp_path / "origen"
    origen.mkdir()
    for nombre in (
        "INV-2026-0001 factura.pdf",
        "INV-2026-0001 anexo.pdf",
        "2026EE963171O1 radicado.pdf",
    ):
        crear_pdf(origen, nombre)

    ruta = tmp_path / "m.xlsx"
    crear_excel(
        ruta,
        [
            ("INV-2026-0001 contrato", "EXP-2026-001"),
            ("2026EE963171O1", "9900000009"),
        ],
    )

    core.CODIGO_FIJO = "ZREE"
    regs = core.inventariar(origen, modo=core.MODO_AUTO)
    mapeo, conflictos, _informe = core.cargar_mapeo_detallado(
        ruta, modo=core.MODO_AUTO
    )
    plan = core.construir_plan(regs, mapeo, "{id}-{c}-{n}", conflictos)
    salida = {p["archivo"]: p["destino"] for p in plan}

    assert salida["INV-2026-0001 factura.pdf"] == "EXP-2026-001-ZREE-2.pdf"
    assert salida["INV-2026-0001 anexo.pdf"] == "EXP-2026-001-ZREE-1.pdf"
    assert salida["2026EE963171O1 radicado.pdf"] == "9900000009-ZREE-1.pdf"

    permitido, _avisos = core.validar_plan(plan, tmp_path / "destino")
    assert permitido
    copiados, omitidos = core.ejecutar_copias(plan, origen, tmp_path / "destino")
    assert copiados == 3 and omitidos == 0


# --- Modo personalizado -----------------------------------------------------------


def test_modo_custom_regex(tmp_path):
    origen = tmp_path / "origen"
    origen.mkdir()
    crear_pdf(origen, "EXP-2026-001 acta.pdf")
    crear_pdf(origen, "EXP-2026-002 acta.pdf")

    ruta = tmp_path / "m.xlsx"
    crear_excel(
        ruta,
        [
            ("EXP-2026-001", "CT-2026-A"),
            ("EXP-2026-002", "CT-2026-B"),
        ],
    )

    regex_origen = core.re.compile(r"EXP-\d{4}-\d{3}")
    regex_destino = core.re.compile(r"CT-\d{4}-[A-Z]")

    core.CODIGO_FIJO = "V2"
    regs = core.inventariar(
        origen, modo=core.MODO_CUSTOM, regex_origen=regex_origen
    )
    assert [r["id_origen"] for r in regs] == ["EXP-2026-001", "EXP-2026-002"]

    mapeo, _conflictos, _informe = core.cargar_mapeo_detallado(
        ruta, modo=core.MODO_CUSTOM, regex_origen=regex_origen,
        regex_destino=regex_destino,
    )
    assert mapeo == {"EXP-2026-001": "CT-2026-A", "EXP-2026-002": "CT-2026-B"}

    plan = core.construir_plan(regs, mapeo, "{id}-{c}-{n}")
    assert plan[0]["destino"] == "CT-2026-A-V2-1.pdf"


def test_modo_custom_regex_invalida(tmp_path):
    origen = tmp_path / "origen"
    origen.mkdir()

    with pytest.raises(ValueError):
        core.inventariar(
            origen, modo=core.MODO_CUSTOM, regex_origen="(EXP-"
        )
    with pytest.raises(ValueError):
        core.cargar_mapeo_detallado(
            tmp_path / "no_importa.xlsx", modo=core.MODO_CUSTOM
        )


def test_plantilla_excel_generada(tmp_path):
    ruta = tmp_path / "plantilla_correspondencia.xlsx"
    core.crear_plantilla_excel(ruta)

    assert ruta.is_file()
    wb = openpyxl.load_workbook(ruta)
    assert wb.sheetnames == ["Correspondencia", "Instrucciones"]
    ws = wb["Correspondencia"]
    filas = [r for r in ws.iter_rows(values_only=True)]
    assert filas[0] == ("ID origen", "ID destino")
    assert len(filas) == 4

    mapeo, conflictos, informe = core.cargar_mapeo_detallado(
        ruta, modo=core.MODO_AUTO
    )
    assert informe["pares"] >= 2
    assert "2026XX840820X1" in mapeo or "INV-2026-0001" in mapeo

    # metadata sin rastros de herramientas
    import zipfile

    with zipfile.ZipFile(ruta) as z:
        app = z.read("docProps/app.xml").decode("utf-8", "replace")
        core_xml = z.read("docProps/core.xml").decode("utf-8", "replace")
    assert "openpyxl" not in app.lower()
    assert "openpyxl" not in core_xml.lower()